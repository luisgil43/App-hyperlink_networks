import json
import os
import re
from datetime import timedelta
from decimal import Decimal
from io import BytesIO

from django.conf import settings
from django.contrib import messages
from django.contrib.admin.views.decorators import staff_member_required
from django.contrib.auth import authenticate, get_user_model, login, logout
from django.contrib.auth.decorators import login_required
from django.contrib.auth.forms import AuthenticationForm
from django.contrib.auth.models import Group
from django.contrib.auth.views import LoginView
from django.contrib.staticfiles import finders
from django.core.exceptions import PermissionDenied
from django.core.paginator import EmptyPage, PageNotAnInteger, Paginator
from django.db.models import Count, Prefetch, Q, Sum
from django.http import HttpResponse
from django.shortcuts import get_object_or_404, redirect, render
from django.urls import NoReverseMatch, reverse, reverse_lazy
from django.utils import timezone
from django.utils.http import url_has_allowed_host_and_scheme
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement  # 👈 para cantSplit
from docx.shared import Inches, Pt, RGBColor

from core.permissions import filter_queryset_by_access
from dashboard.models import ProduccionTecnico
from facturacion.models import Proyecto
from operaciones.models import SesionBilling
from rrhh.forms import FeriadoForm
from rrhh.models import Feriado
from usuarios.decoradores import rol_requerido
from usuarios.models import CustomUser
from usuarios.models import CustomUser as User
from usuarios.models import Notificacion, Rol, TrustedDevice, get_role_label_en

# Intentamos ubicar el modelo de asignación (si ya existe con through)
try:
    from usuarios.models import ProyectoAsignacion  # through recomendado
except Exception:
    try:
        from facturacion.models import ProyectoAsignacion
    except Exception:
        ProyectoAsignacion = None  # fallback si aún no creas el through
User = get_user_model()


@login_required(login_url='usuarios:login')
def admin_dashboard_view(request):
    # Cargar datos para la plantilla principal del admin dashboard
    return render(request, 'dashboard_admin/base.html')


@login_required(login_url='usuarios:login_unificado')
def logout_view(request):
    logout(request)
    messages.info(request, "You have successfully logged out.")
    return redirect(reverse('usuarios:login_unificado'))


@login_required(login_url="usuarios:login_unificado")
@rol_requerido("admin", "pm", "supervisor", "facturacion")
def inicio_admin(request):
    user = request.user
    today = timezone.localdate()

    start_week = today - timedelta(days=today.weekday())
    end_week = start_week + timedelta(days=6)

    start_prev_week = start_week - timedelta(days=7)
    end_prev_week = start_week - timedelta(days=1)

    iso_year, iso_week, _ = today.isocalendar()
    week_label = f"{iso_year}-W{int(iso_week):02d}"
    week_range_label = (
        f"{start_week.strftime('%b %d')} - {end_week.strftime('%b %d, %Y')}"
    )

    qs = SesionBilling.objects.all()

    # ======================================================
    # FILTRO POR PROYECTOS VISIBLES
    # Admin/superuser ve todo.
    # PM/supervisor/facturación ve solo proyectos asignados.
    # ======================================================
    can_view_all = user.is_superuser or getattr(user, "es_admin_general", False)

    if not can_view_all:
        proyectos_user = filter_queryset_by_access(
            Proyecto.objects.all(),
            user,
            "id",
        )

        proyectos_user_list = list(proyectos_user)

        if proyectos_user_list:
            allowed_keys = set()

            for p in proyectos_user_list:
                nombre = (getattr(p, "nombre", "") or "").strip()
                codigo = (getattr(p, "codigo", "") or "").strip()

                if nombre:
                    allowed_keys.add(nombre)

                if codigo:
                    allowed_keys.add(codigo)

                allowed_keys.add(str(p.id))

            qs = qs.filter(
                Q(proyecto__in=allowed_keys) | Q(proyecto_id__in=allowed_keys)
            )
        else:
            qs = qs.none()

    approved_states = ["aprobado_supervisor", "aprobado_pm"]

    total_assigned = qs.filter(estado="asignado").count()
    total_in_progress = qs.filter(estado="en_proceso").count()
    total_submitted_review = qs.filter(estado="en_revision_supervisor").count()

    approved_week = qs.filter(
        estado__in=approved_states,
        creado_en__date__range=[start_week, end_week],
    ).count()

    approved_prev_week = qs.filter(
        estado__in=approved_states,
        creado_en__date__range=[start_prev_week, end_prev_week],
    ).count()

    total_current = (
        total_assigned + total_in_progress + total_submitted_review + approved_week
    )

    performance = round((approved_week / total_current) * 100) if total_current else 0

    if approved_prev_week > 0:
        vs_last_week = round(
            ((approved_week - approved_prev_week) / approved_prev_week) * 100
        )
    else:
        vs_last_week = 100 if approved_week > 0 else 0

    chart_labels = ["Mon", "Tue", "Wed", "Thu", "Fri", "Sat", "Sun"]
    chart_data = []

    for i in range(7):
        day = start_week + timedelta(days=i)

        chart_data.append(
            qs.filter(
                estado__in=approved_states,
                creado_en__date=day,
            ).count()
        )

    queryset = Notificacion.objects.filter(usuario=user).order_by("leido", "-fecha")

    return render(
        request,
        "dashboard_admin/inicio_admin.html",
        {
            "notificaciones": queryset[:10],
            "notificaciones_no_leidas": queryset.filter(leido=False).count(),
            "week_label": week_label,
            "week_range_label": week_range_label,
            "total_assigned": total_assigned,
            "total_in_progress": total_in_progress,
            "total_submitted_review": total_submitted_review,
            "approved_week": approved_week,
            "approved_prev_week": approved_prev_week,
            "performance": performance,
            "vs_last_week": vs_last_week,
            "chart_labels": json.dumps(chart_labels),
            "chart_data": json.dumps(chart_data),
        },
    )


@login_required(login_url='usuarios:login')
@rol_requerido('admin', 'pm', 'supervisor')
def produccion_tecnico(request):
    produccion = ProduccionTecnico.objects.filter(tecnico__user=request.user)
    return render(request, 'dashboard/produccion_tecnico.html', {
        'produccion': produccion
    })


@login_required(login_url='usuarios:login')
@rol_requerido('admin', 'pm', 'rrhh')
def grupos_view(request):
    if request.method == 'POST':
        nombre = request.POST.get('nombre', '').strip()
        grupo_id = request.POST.get('grupo_id')

        if 'add_group' in request.POST:
            if nombre:
                grupo, creado = Group.objects.get_or_create(name=nombre)
                if creado:
                    messages.success(request, f'Group "{nombre}" created successfully.')
                else:
                    messages.warning(request, f'Group "{nombre}" already exists.')
            else:
                messages.error(request, "You must enter a group name.")
            return redirect('dashboard_admin:grupos')

        elif 'delete_group' in request.POST and grupo_id:
            try:
                grupo = Group.objects.get(id=grupo_id)
                grupo.delete()
                messages.success(request, f'Group "{grupo.name}" deleted successfully.')
            except Group.DoesNotExist:
                messages.error(request, 'Group does not exist.')
            return redirect('dashboard_admin:grupos')

    grupos = Group.objects.all().order_by('name')
    return render(request, 'dashboard_admin/grupos.html', {'grupos': grupos})


@login_required(login_url="usuarios:login")
@rol_requerido("admin", "pm")
def editar_usuario_view(request, user_id):
    from django.db import transaction
    from django.utils import timezone

    usuario = get_object_or_404(User, id=user_id)
    grupos = Group.objects.all()
    roles_disponibles = Rol.objects.all()

    # ==============================================================
    # SEGURIDAD
    # ==============================================================

    actor_is_admin = (
        request.user.is_superuser
        or request.user.roles.filter(nombre__iexact="admin").exists()
    )

    target_has_admin_role = usuario.roles.filter(nombre__iexact="admin").exists()

    target_is_privileged = (
        usuario.is_superuser or usuario.is_staff or target_has_admin_role
    )

    # Un usuario no-admin no puede administrar cuentas privilegiadas.
    if not actor_is_admin and target_is_privileged:
        messages.error(
            request,
            "You do not have permission to modify this privileged account.",
        )

        return redirect("dashboard_admin:listar_usuarios")

    # ==============================================================
    # DESBLOQUEO DE EDICIÓN
    #
    # La palabra se solicita ANTES de mostrar el formulario.
    # Se guarda únicamente un desbloqueo temporal en sesión.
    # La palabra nunca se guarda.
    # ==============================================================

    session_key = f"user_management_edit_unlock_{usuario.id}"

    unlock_data = request.session.get(session_key)

    edit_unlocked = False

    if isinstance(
        unlock_data,
        dict,
    ):
        try:
            unlocked_at = float(
                unlock_data.get(
                    "timestamp",
                    0,
                )
            )

            unlocked_by = int(
                unlock_data.get(
                    "actor_id",
                    0,
                )
            )

            elapsed = timezone.now().timestamp() - unlocked_at

            edit_unlocked = unlocked_by == request.user.id and elapsed <= 300

        except (
            TypeError,
            ValueError,
        ):
            edit_unlocked = False

    # Si el desbloqueo expiró o es inválido, lo limpiamos.
    if unlock_data and not edit_unlocked:
        request.session.pop(
            session_key,
            None,
        )

        request.session.modified = True

    # ==============================================================
    # POST ESPECIAL:
    # desbloquear edición antes de entrar al formulario.
    # ==============================================================

    if request.method == "POST" and "unlock_user_edit" in request.POST:
        confirmation_word = request.POST.get("management_confirmation_word") or ""

        if not _valid_user_management_confirmation(confirmation_word):
            messages.error(
                request,
                "Invalid security confirmation word. "
                "Access to user editing was denied.",
            )

            return redirect("dashboard_admin:listar_usuarios")

        request.session[session_key] = {
            "actor_id": request.user.id,
            "timestamp": timezone.now().timestamp(),
        }

        request.session.modified = True

        return redirect(
            "dashboard_admin:editar_usuario",
            user_id=usuario.id,
        )

    # ==============================================================
    # SIN DESBLOQUEO:
    # no mostrar formulario.
    # ==============================================================

    if not edit_unlocked:
        messages.error(
            request,
            "Security confirmation is required " "before editing this user.",
        )

        return redirect("dashboard_admin:listar_usuarios")

    # ==============================================================
    # POST NORMAL:
    # guardar edición.
    #
    # IMPORTANTE:
    # Esta vista NO administra proyectos.
    # ProyectoAsignacion no se modifica aquí.
    # ==============================================================

    if request.method == "POST":
        username = (request.POST.get("username") or usuario.username).strip()

        # ==========================================================
        # Username reservado.
        # ==========================================================

        if username.lower() == "admin":
            messages.error(
                request,
                'The username "admin" is reserved ' "and cannot be used.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        # ==========================================================
        # Username único sin importar mayúsculas/minúsculas.
        # ==========================================================

        if (
            User.objects.filter(username__iexact=username)
            .exclude(id=usuario.id)
            .exists()
        ):
            messages.error(
                request,
                "Username already exists.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        # ==========================================================
        # SUPERUSER:
        # puede editarse, pero su estado no puede cambiar.
        # ==========================================================

        requested_is_superuser = "is_superuser" in request.POST

        if requested_is_superuser != usuario.is_superuser:
            messages.error(
                request,
                "Superuser status cannot be changed " "from Planix.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        # ==========================================================
        # Roles.
        # ==========================================================

        roles_ids = request.POST.getlist("roles")

        roles_solicitados = list(Rol.objects.filter(id__in=roles_ids))

        submitted_role_ids = {str(rol.id) for rol in roles_solicitados}

        requested_role_ids = {str(role_id) for role_id in roles_ids}

        if submitted_role_ids != requested_role_ids:
            messages.error(
                request,
                "One or more selected roles are invalid.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        if not roles_solicitados:
            messages.error(
                request,
                "Please select at least one role.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        requested_admin_role = any(
            (rol.nombre or "").strip().lower() == "admin" for rol in roles_solicitados
        )

        if requested_admin_role and not actor_is_admin:
            messages.error(
                request,
                "You do not have permission " "to assign the Admin role.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        # ==========================================================
        # Staff solamente Admin.
        # ==========================================================

        requested_is_staff = "is_staff" in request.POST

        if not actor_is_admin and requested_is_staff != usuario.is_staff:
            messages.error(
                request,
                "You do not have permission " "to modify Staff privileges.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        # ==========================================================
        # Password.
        #
        # EDIT:
        # ambos vacíos = conservar password actual.
        # Si se escribe alguno, ambos deben existir y coincidir.
        # ==========================================================

        password1 = request.POST.get("password1") or ""

        password2 = request.POST.get("password2") or ""

        if password1 or password2:
            if not password1:
                messages.error(
                    request,
                    "Please enter the new password.",
                )

                return redirect(
                    "dashboard_admin:editar_usuario",
                    user_id=usuario.id,
                )

            if not password2:
                messages.error(
                    request,
                    "Please confirm the password.",
                )

                return redirect(
                    "dashboard_admin:editar_usuario",
                    user_id=usuario.id,
                )

            if password1 != password2:
                messages.error(
                    request,
                    "Passwords do not match.",
                )

                return redirect(
                    "dashboard_admin:editar_usuario",
                    user_id=usuario.id,
                )

        # ==========================================================
        # Datos POST.
        # ==========================================================

        grupo_ids = request.POST.getlist("groups")

        identidad_post = (request.POST.get("identidad") or "").strip()

        if identidad_post and not re.match(
            r"^[A-Za-z0-9\.\-]+$",
            identidad_post,
        ):
            messages.error(
                request,
                "ID may contain only letters, numbers, " "dots, or hyphens.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        if (
            identidad_post
            and User.objects.filter(identidad=identidad_post)
            .exclude(id=usuario.id)
            .exists()
        ):
            messages.error(
                request,
                "ID number is already registered.",
            )

            return redirect(
                "dashboard_admin:editar_usuario",
                user_id=usuario.id,
            )

        # ==========================================================
        # Escritura atómica.
        #
        # NO SE TOCAN PROYECTOS.
        # ==============================================================

        with transaction.atomic():
            usuario.username = username

            usuario.first_name = request.POST.get(
                "first_name",
                usuario.first_name,
            )

            usuario.last_name = request.POST.get(
                "last_name",
                usuario.last_name,
            )

            usuario.email = request.POST.get(
                "email",
                usuario.email,
            )

            usuario.is_active = "is_active" in request.POST

            if actor_is_admin:
                usuario.is_staff = requested_is_staff

            # is_superuser NO se modifica.
            usuario.identidad = identidad_post

            if password1:
                usuario.set_password(password1)

            usuario.save()

            usuario.groups.set(grupo_ids)

            usuario.roles.set(roles_ids)

            # ======================================================
            # IMPORTANTE:
            # ProyectoAsignacion NO se modifica aquí.
            #
            # Todas las asignaciones existentes permanecen intactas.
            # La administración de proyectos se realiza únicamente
            # desde gestionar_asignaciones_proyectos_view.
            # ======================================================

        # ==========================================================
        # El desbloqueo es de un solo uso después de guardar.
        # ==========================================================

        request.session.pop(
            session_key,
            None,
        )

        request.session.modified = True

        messages.success(
            request,
            "User updated successfully.",
        )

        return redirect("dashboard_admin:listar_usuarios")

    # ==============================================================
    # GET
    # ==============================================================

    roles_seleccionados = [
        str(pk)
        for pk in usuario.roles.values_list(
            "id",
            flat=True,
        )
    ]

    grupo_ids_post = [
        str(gid)
        for gid in usuario.groups.values_list(
            "id",
            flat=True,
        )
    ]

    return render(
        request,
        "dashboard_admin/editar_usuario.html",
        {
            "usuario": usuario,
            "grupos": grupos,
            "roles": roles_disponibles,
            "roles_seleccionados": roles_seleccionados,
            "grupo_ids_post": grupo_ids_post,
        },
    )


def _valid_user_management_confirmation(value):
    import hmac
    import os

    expected = (
        os.environ.get(
            "USER_MANAGEMENT_CONFIRMATION_WORD",
            "",
        )
        or ""
    )

    received = value or ""

    # Fail closed:
    # si la variable no está configurada, nunca autoriza.
    if not expected:
        return False

    return hmac.compare_digest(
        received,
        expected,
    )


@login_required(login_url="usuarios:login")
@rol_requerido("admin", "pm", "rrhh")
def crear_usuario_view(request, identidad=None):
    from django.db import transaction

    grupos = Group.objects.all()

    usuario = (
        get_object_or_404(
            User,
            identidad=identidad,
        )
        if identidad
        else None
    )

    # ==============================================================
    # SEGURIDAD
    # ==============================================================

    actor_is_admin = (
        request.user.is_superuser
        or request.user.roles.filter(nombre__iexact="admin").exists()
    )

    if usuario:
        target_has_admin_role = usuario.roles.filter(nombre__iexact="admin").exists()

        target_is_privileged = (
            usuario.is_superuser or usuario.is_staff or target_has_admin_role
        )

        if not actor_is_admin and target_is_privileged:
            messages.error(
                request,
                "You do not have permission " "to modify this privileged account.",
            )

            return redirect("dashboard_admin:listar_usuarios")

        if usuario.is_superuser:
            messages.error(
                request,
                "Superuser accounts cannot be modified "
                "from this user creation view.",
            )

            return redirect("dashboard_admin:listar_usuarios")

    if request.method == "POST":
        username = (request.POST.get("username") or "").strip()

        email = (request.POST.get("email") or "").strip()

        password1 = request.POST.get("password1") or ""

        password2 = request.POST.get("password2") or ""

        first_name = (request.POST.get("first_name") or "").strip()

        last_name = (request.POST.get("last_name") or "").strip()

        is_active = request.POST.get("is_active") == "on"

        requested_is_staff = "is_staff" in request.POST

        grupo_ids_raw = request.POST.getlist("groups")

        try:
            grupo_ids = [int(gid) for gid in grupo_ids_raw]

        except (
            TypeError,
            ValueError,
        ):
            messages.error(
                request,
                "One or more selected groups are invalid.",
            )

            return redirect(request.path)

        identidad_post = (request.POST.get("identidad") or "").strip()

        roles_ids = request.POST.getlist("roles")

        # ==========================================================
        # Campos obligatorios.
        # ==========================================================

        if not username:
            messages.error(
                request,
                "Username is required.",
            )

            return redirect(request.path)

        if not email:
            messages.error(
                request,
                "Email is required.",
            )

            return redirect(request.path)

        if not first_name:
            messages.error(
                request,
                "First name is required.",
            )

            return redirect(request.path)

        if not last_name:
            messages.error(
                request,
                "Last name is required.",
            )

            return redirect(request.path)

        if not identidad_post:
            messages.error(
                request,
                "ID / Identification Number is required.",
            )

            return redirect(request.path)

        # ==========================================================
        # Username reservado.
        # ==========================================================

        if username.lower() == "admin":
            messages.error(
                request,
                'The username "admin" is reserved ' "and cannot be used.",
            )

            return redirect(request.path)

        # ==========================================================
        # is_superuser nunca puede asignarse desde Planix.
        # ==========================================================

        if "is_superuser" in request.POST:
            messages.error(
                request,
                "Superuser privileges cannot be assigned " "from Planix.",
            )

            return redirect(request.path)

        # ==========================================================
        # Validación de roles.
        # ==========================================================

        roles_solicitados = list(Rol.objects.filter(id__in=roles_ids))

        submitted_role_ids = {str(rol.id) for rol in roles_solicitados}

        requested_role_ids = {str(role_id) for role_id in roles_ids}

        if submitted_role_ids != requested_role_ids:
            messages.error(
                request,
                "One or more selected roles are invalid.",
            )

            return redirect(request.path)

        if not roles_solicitados:
            messages.error(
                request,
                "Please select at least one role.",
            )

            return redirect(request.path)

        requested_admin_role = any(
            (rol.nombre or "").strip().lower() == "admin" for rol in roles_solicitados
        )

        if requested_admin_role and not actor_is_admin:
            messages.error(
                request,
                "You do not have permission " "to assign the Admin role.",
            )

            return redirect(request.path)

        # ==========================================================
        # is_staff solamente Admin.
        # ==========================================================

        if requested_is_staff and not actor_is_admin:
            messages.error(
                request,
                "You do not have permission " "to assign Staff privileges.",
            )

            return redirect(request.path)

        # ==========================================================
        # Campos jerárquicos.
        # ==============================================================

        def get_user_or_none(uid):
            return CustomUser.objects.filter(id=uid).first() if uid else None

        supervisor = get_user_or_none(request.POST.get("supervisor"))

        pm = get_user_or_none(request.POST.get("pm"))

        rrhh_encargado = get_user_or_none(request.POST.get("rrhh_encargado"))

        prevencionista = get_user_or_none(request.POST.get("prevencionista"))

        logistica_encargado = get_user_or_none(request.POST.get("logistica_encargado"))

        encargado_flota = get_user_or_none(request.POST.get("encargado_flota"))

        encargado_subcontrato = get_user_or_none(
            request.POST.get("encargado_subcontrato")
        )

        encargado_facturacion = get_user_or_none(
            request.POST.get("encargado_facturacion")
        )

        # ==========================================================
        # Password.
        #
        # CREATE:
        # obligatorio.
        #
        # EDIT DESDE ESTA VISTA LEGACY:
        # opcional.
        # ==============================================================

        if usuario is None:
            if not password1:
                messages.error(
                    request,
                    "Password is required to create a user.",
                )

                return redirect(request.path)

            if not password2:
                messages.error(
                    request,
                    "Please confirm the password.",
                )

                return redirect(request.path)

            if password1 != password2:
                messages.error(
                    request,
                    "Passwords do not match.",
                )

                return redirect(request.path)

        else:
            if password1 or password2:
                if not password1:
                    messages.error(
                        request,
                        "Please enter the new password.",
                    )

                    return redirect(request.path)

                if not password2:
                    messages.error(
                        request,
                        "Please confirm the password.",
                    )

                    return redirect(request.path)

                if password1 != password2:
                    messages.error(
                        request,
                        "Passwords do not match.",
                    )

                    return redirect(request.path)

        # ==========================================================
        # Validación ID.
        # ==============================================================

        if identidad_post and not re.match(
            r"^[A-Za-z0-9\.\-]+$",
            identidad_post,
        ):
            messages.error(
                request,
                "ID may contain only letters, numbers, " "dots, or hyphens.",
            )

            return redirect(request.path)

        # ==========================================================
        # Edición desde esta misma vista.
        #
        # Se conserva por compatibilidad con identidad=None/identity.
        # NO administra proyectos.
        # ==============================================================

        if usuario:
            if (
                User.objects.filter(username__iexact=username)
                .exclude(id=usuario.id)
                .exists()
            ):
                messages.error(
                    request,
                    "Username already exists.",
                )

                return redirect(request.path)

            if (
                identidad_post
                and User.objects.filter(identidad=identidad_post)
                .exclude(id=usuario.id)
                .exists()
            ):
                messages.error(
                    request,
                    "ID number is already registered.",
                )

                return redirect(request.path)

            with transaction.atomic():
                usuario.username = username
                usuario.email = email
                usuario.first_name = first_name
                usuario.last_name = last_name
                usuario.is_active = is_active

                if actor_is_admin:
                    usuario.is_staff = requested_is_staff

                # IMPORTANTE:
                # is_superuser no se toca.
                usuario.identidad = identidad_post

                usuario.supervisor = supervisor
                usuario.pm = pm

                usuario.rrhh_encargado = rrhh_encargado

                usuario.prevencionista = prevencionista

                usuario.logistica_encargado = logistica_encargado

                usuario.encargado_flota = encargado_flota

                usuario.encargado_subcontrato = encargado_subcontrato

                usuario.encargado_facturacion = encargado_facturacion

                if password1:
                    usuario.set_password(password1)

                usuario.save()

                usuario.groups.set(grupo_ids)

                usuario.roles.set(roles_ids)

                # ==================================================
                # ProyectoAsignacion NO se modifica aquí.
                # ==================================================

            messages.success(
                request,
                "User updated successfully.",
            )

        else:
            # ======================================================
            # Creación.
            # ======================================================

            if User.objects.filter(username__iexact=username).exists():
                messages.error(
                    request,
                    "Username already exists.",
                )

                return redirect("dashboard_admin:crear_usuario")

            if (
                identidad_post
                and User.objects.filter(identidad=identidad_post).exists()
            ):
                messages.error(
                    request,
                    "ID number is already registered.",
                )

                return redirect("dashboard_admin:crear_usuario")

            with transaction.atomic():
                usuario = User.objects.create_user(
                    username=username,
                    email=email,
                    password=password1,
                    first_name=first_name,
                    last_name=last_name,
                    is_active=is_active,
                    # Staff únicamente si quien crea es Admin.
                    is_staff=(requested_is_staff if actor_is_admin else False),
                    # NUNCA desde Planix.
                    is_superuser=False,
                    identidad=identidad_post,
                    supervisor=supervisor,
                    pm=pm,
                    rrhh_encargado=(rrhh_encargado),
                    prevencionista=(prevencionista),
                    logistica_encargado=(logistica_encargado),
                    encargado_flota=(encargado_flota),
                    encargado_subcontrato=(encargado_subcontrato),
                    encargado_facturacion=(encargado_facturacion),
                )

                usuario.groups.set(grupo_ids)

                usuario.roles.set(roles_ids)

                # ==================================================
                # IMPORTANTE:
                # El usuario se crea SIN proyectos.
                #
                # Las asignaciones se realizan exclusivamente desde:
                # Project Assignments.
                # ==================================================

            messages.success(
                request,
                "User created successfully.",
            )

        return redirect("dashboard_admin:listar_usuarios")

    # ==============================================================
    # GET
    # ==============================================================

    grupo_ids_post = request.POST.getlist("groups") if request.method == "POST" else []

    if not grupo_ids_post and usuario:
        grupo_ids_post = [str(g.id) for g in usuario.groups.all()]

    roles_disponibles = Rol.objects.all()

    roles_seleccionados = (
        usuario.roles.values_list(
            "id",
            flat=True,
        )
        if usuario
        else []
    )

    roles_seleccionados = [str(role_id) for role_id in roles_seleccionados]

    usuarios_activos = CustomUser.objects.filter(is_active=True).order_by(
        "first_name",
        "last_name",
    )

    contexto = {
        "grupos": grupos,
        "grupo_ids_post": grupo_ids_post,
        "usuario": usuario,
        "roles": roles_disponibles,
        "roles_seleccionados": roles_seleccionados,
        "usuarios": usuarios_activos,
    }

    return render(
        request,
        "dashboard_admin/crear_usuario.html",
        contexto,
    )


@login_required(login_url="usuarios:login")
@rol_requerido("admin", "pm", "rrhh")
def listar_usuarios(request):
    # ==============================================================
    # ¿El usuario actual puede eliminar cuentas?
    #
    # Delete es exclusivo del rol Admin.
    # is_superuser también se considera administrador del sistema.
    # ==============================================================

    can_delete_users = (
        request.user.is_superuser
        or request.user.roles.filter(nombre__iexact="admin").exists()
    )

    # ==============================================================
    # ACCIONES POST
    # ==============================================================

    if request.method == "POST":
        user_id = request.POST.get("user_id")

        # ----------------------------------------------------------
        # 1) ACTIVAR / DESACTIVAR USUARIO
        # ----------------------------------------------------------

        if "toggle_user_status" in request.POST:
            try:
                usuario = User.objects.get(id=user_id)
            except User.DoesNotExist:
                messages.error(request, "User not found.")
                return redirect("dashboard_admin:listar_usuarios")

            # Nadie puede desactivarse a sí mismo desde esta pantalla.
            if usuario.id == request.user.id:
                messages.error(
                    request,
                    "You cannot activate or deactivate your own account from this screen.",
                )
                return redirect("dashboard_admin:listar_usuarios")

            action = (request.POST.get("user_status_action") or "").strip().lower()

            if action == "deactivate":
                usuario.is_active = False
                usuario.save(update_fields=["is_active"])

                messages.success(
                    request, f'User "{usuario.username}" has been deactivated.'
                )

            elif action == "activate":
                usuario.is_active = True
                usuario.save(update_fields=["is_active"])

                messages.success(
                    request, f'User "{usuario.username}" has been activated.'
                )

            else:
                messages.error(request, "Invalid user status action.")

            return redirect("dashboard_admin:listar_usuarios")

        # ----------------------------------------------------------
        # 2) RESET 2FA
        # ----------------------------------------------------------

        if "reset_2fa" in request.POST:
            try:
                usuario = User.objects.get(id=user_id)
            except User.DoesNotExist:
                messages.error(request, "User not found.")
                return redirect("dashboard_admin:listar_usuarios")

            update_fields = []

            if hasattr(
                usuario,
                "two_factor_secret",
            ):
                usuario.two_factor_secret = ""
                update_fields.append("two_factor_secret")

            if hasattr(
                usuario,
                "two_factor_enabled",
            ):
                usuario.two_factor_enabled = False
                update_fields.append("two_factor_enabled")

            if update_fields:
                usuario.save(update_fields=update_fields)
            else:
                usuario.save()

            try:
                TrustedDevice.objects.filter(user=usuario).delete()
            except Exception:
                pass

            messages.success(
                request,
                f'2FA has been reset for user "{usuario.username}". '
                "They must configure it again.",
            )

            return redirect("dashboard_admin:listar_usuarios")

    # ==============================================================
    # FILTROS GET
    # ==============================================================

    rol_filtrado = (request.GET.get("rol") or "").strip()

    first_q = (request.GET.get("first") or "").strip()

    last_q = (request.GET.get("last") or "").strip()

    id_q = (request.GET.get("id") or "").strip()

    qs = (
        User.objects.all()
        .order_by("id")
        .prefetch_related(
            "roles",
            "groups",
        )
    )

    # Prefetch de proyectos según exista through o M2M directo
    if ProyectoAsignacion:
        qs = qs.prefetch_related(
            Prefetch(
                "proyectoasignacion_set",
                queryset=(ProyectoAsignacion.objects.select_related("proyecto")),
            )
        )

    elif hasattr(User, "proyectos"):
        qs = qs.prefetch_related("proyectos")

    if rol_filtrado:
        qs = qs.filter(roles__nombre=rol_filtrado).distinct()

    if first_q:
        qs = qs.filter(first_name__icontains=first_q)

    if last_q:
        qs = qs.filter(last_name__icontains=last_q)

    if id_q:
        qs = qs.filter(identidad__icontains=id_q)

    # ==============================================================
    # PAGINACIÓN
    # ==============================================================

    per_page_raw = (
        str(
            request.GET.get(
                "per_page",
                request.GET.get(
                    "cantidad",
                    "20",
                ),
            )
        )
        .strip()
        .lower()
    )

    if per_page_raw in (
        "all",
        "todos",
    ):
        per_page = max(
            qs.count(),
            1,
        )

    else:
        try:
            per_page = int(per_page_raw or 20)
        except ValueError:
            per_page = 20

        per_page = max(
            5,
            min(
                per_page,
                100,
            ),
        )

    paginator = Paginator(
        qs,
        per_page,
    )

    page_number = request.GET.get(
        "page",
        1,
    )

    try:
        usuarios_page = paginator.get_page(page_number)

    except (
        PageNotAnInteger,
        EmptyPage,
    ):
        usuarios_page = paginator.get_page(1)

    # Preserva querystring excepto page
    params = request.GET.copy()
    params.pop(
        "page",
        None,
    )

    querystring = params.urlencode()

    roles_disponibles = Rol.objects.all()

    return render(
        request,
        "dashboard_admin/listar_usuarios.html",
        {
            "usuarios": usuarios_page,
            "page_obj": usuarios_page,
            "roles": roles_disponibles,
            "rol_filtrado": rol_filtrado,
            "per_page": per_page,
            "querystring": querystring,
            "first_q": first_q,
            "last_q": last_q,
            "id_q": id_q,
            "cantidad": request.GET.get(
                "cantidad",
                None,
            ),
            # Hace funcionar la condición del HTML.
            "can_delete_users": can_delete_users,
        },
    )


@login_required(login_url="usuarios:login")
@rol_requerido("admin")
def eliminar_usuario_view(request, user_id):
    from django.db.models.deletion import ProtectedError

    usuario = get_object_or_404(
        User,
        id=user_id,
    )

    # ==============================================================
    # BLINDAJE 1:
    # Nunca permitir borrar la cuenta reservada "admin".
    # ==============================================================

    if (usuario.username or "").strip().lower() == "admin":
        messages.error(request, 'The reserved "admin" account cannot be deleted.')

        return redirect("dashboard_admin:listar_usuarios")

    # ==============================================================
    # BLINDAJE 2:
    # Nunca permitir que un usuario se elimine a sí mismo.
    # ==============================================================

    if usuario.id == request.user.id:
        messages.error(request, "You cannot delete your own account.")

        return redirect("dashboard_admin:listar_usuarios")

    # ==============================================================
    # DELETE
    # ==============================================================

    if request.method == "POST":
        confirmation_word = request.POST.get("management_confirmation_word") or ""

        # ==========================================================
        # BLINDAJE 3:
        # Confirmación privada del lado servidor.
        # ==========================================================

        if not _valid_user_management_confirmation(confirmation_word):
            messages.error(
                request,
                "Invalid security confirmation word. " "The user was not deleted.",
            )

            return redirect("dashboard_admin:listar_usuarios")

        username = usuario.username

        try:
            usuario.delete()

        except ProtectedError:
            messages.error(
                request,
                f'User "{username}" cannot be deleted because '
                "historical records are still linked to this account. "
                "Deactivate the user instead or reassign the protected "
                "historical records first.",
            )

            return redirect("dashboard_admin:listar_usuarios")

        messages.success(request, f'User "{username}" deleted successfully.')

        return redirect("dashboard_admin:listar_usuarios")

    return render(
        request,
        "dashboard_admin/eliminar_usuario_confirmacion.html",
        {
            "usuario": usuario,
        },
    )


# Vista para usuarios no autorizados
def no_autorizado(request):
    return render(request, 'dashboard_admin/no_autorizado.html')


@login_required
def redireccionar_vacaciones(request):
    user = request.user
    if user.es_supervisor:
        return redirect('rrhh:revisar_supervisor')
    elif user.es_pm:
        return redirect('rrhh:revisar_pm')
    elif user.es_rrhh or user.es_admin_general:  # 👈 Aquí
        return redirect('rrhh:revisar_rrhh')
    else:
        return redirect('dashboard_admin:inicio_admin')


@login_required
@rol_requerido('rrhh')
def listar_feriados(request):
    feriados = Feriado.objects.order_by('fecha')
    form = FeriadoForm()

    if request.method == 'POST':
        form = FeriadoForm(request.POST)
        if form.is_valid():
            form.save()
            return redirect('dashboard_admin:listar_feriados')

    return render(request, 'dashboard_admin/listar_feriados.html', {
        'feriados': feriados,
        'form': form
    })


@login_required
@rol_requerido('rrhh')
def eliminar_feriado(request, pk):
    feriado = get_object_or_404(Feriado, pk=pk)
    feriado.delete()
    messages.success(
        request, f'El feriado "{feriado.nombre}" fue eliminado con éxito.')
    return redirect('dashboard_admin:listar_feriados')


def redirigir_a_login_unificado(request):
    return redirect('usuarios:login_unificado')


@login_required(login_url='usuarios:login')
@rol_requerido('admin')
def exportar_usuarios(request):
    """
    Exporta a Excel la lista de usuarios,
    respetando los mismos filtros que listar_usuarios.
    """
    from django.http import HttpResponse
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    # ---- Filtros (mismos names del template) ----
    first_q      = (request.GET.get('first') or '').strip()
    last_q       = (request.GET.get('last') or '').strip()
    id_q         = (request.GET.get('id') or '').strip()
    rol_filtrado = (request.GET.get('rol') or '').strip()

    # Query base
    qs = (
        CustomUser.objects
        .all()
        .prefetch_related(
            'roles',
            'proyectoasignacion_set__proyecto',
            'proyectos',
        )
        .order_by('first_name', 'last_name', 'username')
    )

    if first_q:
        qs = qs.filter(first_name__icontains=first_q)
    if last_q:
        qs = qs.filter(last_name__icontains=last_q)
    if id_q:
        qs = qs.filter(identidad__icontains=id_q)
    if rol_filtrado:
        qs = qs.filter(roles__nombre=rol_filtrado)

    # ---- Crear Excel ----
    wb = Workbook()
    ws = wb.active
    ws.title = "Users"

    headers = [
        "Username",
        "First name",
        "Last name",
        "ID",
        "Email",
        "Active",
        "Staff",
        "Superuser",
        "Roles",
        "Projects",
    ]
    ws.append(headers)

    for u in qs:
        # Roles -> EN INGLÉS usando get_role_label_en
        roles_str = ", ".join(
            get_role_label_en(r.nombre)
            for r in u.roles.all()
        )

        # Proyectos (igual lógica que en el template)
        proyectos_txt = ""

        pas = list(
            u.proyectoasignacion_set.select_related("proyecto").all()
        )
        if pas:
            parts = []
            for a in pas:
                if not a.proyecto:
                    continue
                base = a.proyecto.nombre or ""
                if a.include_history:
                    suf = " (history)"
                elif a.start_at:
                    suf = f" (from {a.start_at.date().isoformat()})"
                else:
                    suf = ""
                parts.append(base + suf)
            proyectos_txt = "; ".join(parts)
        else:
            # Fallback M2M directa u.proyectos
            if hasattr(u, "proyectos"):
                ps = list(u.proyectos.all())
                if ps:
                    proyectos_txt = "; ".join(p.nombre for p in ps)

        ws.append([
            u.username or "",
            u.first_name or "",
            u.last_name or "",
            u.identidad or "",
            u.email or "",
            "Yes" if u.is_active else "No",
            "Yes" if u.is_staff else "No",
            "Yes" if u.is_superuser else "No",
            roles_str or "",
            proyectos_txt or "",
        ])

    # Auto ancho de columnas
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            try:
                max_len = max(max_len, len(str(cell.value or "")))
            except Exception:
                pass
        ws.column_dimensions[col_letter].width = min(max(10, max_len + 2), 50)

    response = HttpResponse(
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        )
    )
    response["Content-Disposition"] = 'attachment; filename="users_export.xlsx"'
    wb.save(response)
    return response


@login_required
@rol_requerido('admin', 'pm', 'facturacion')
def exportar_formato_nuevo_usuario_docx(request):
    """
    Genera un DOCX de solicitud de creación de usuario:
      - Logo (esquina sup. izq.) + títulos centrados en el cuerpo
      - Tabla de datos de usuario
      - Tabla de roles (desde Rol) en INGLÉS
      - Tabla de proyectos (desde Proyecto)
    """

    # Helper: evitar que una fila se parta entre páginas
    def _no_row_split(table):
        for row in table.rows:
            tr = row._tr
            trPr = tr.get_or_add_trPr()
            cant_split = OxmlElement('w:cantSplit')
            trPr.append(cant_split)

    document = Document()

    # ------------------------------------------------------------------
    # LOGO + TÍTULO (fuera del encabezado)
    # ------------------------------------------------------------------
    logo_path = os.path.join(settings.BASE_DIR, "static", "images", "logoh.png")
    logo_par = document.add_paragraph()
    logo_par.alignment = WD_ALIGN_PARAGRAPH.LEFT
    try:
        run_logo = logo_par.add_run()
        run_logo.add_picture(logo_path, width=Inches(1.4))
    except Exception:
        logo_par.add_run("HYPERLINK").bold = True

    title_par = document.add_paragraph()
    title_par.alignment = WD_ALIGN_PARAGRAPH.CENTER

    strong_blue = RGBColor(0x1E, 0x73, 0xBE)  # azul fuerte

    run1 = title_par.add_run("HYPERLINK NETWORKS PLATFORM\n")
    run1.font.name = "Calibri"
    run1.font.size = Pt(13)
    run1.bold = True
    run1.font.color.rgb = strong_blue

    run2 = title_par.add_run("User creation request")
    run2.font.name = "Calibri"
    run2.font.size = Pt(11)
    run2.bold = False
    run2.font.color.rgb = strong_blue

    document.add_paragraph()  # espacio

    # ------------------------------------------------------------------
    # 1) User information
    # ------------------------------------------------------------------
    p_info = document.add_paragraph("User information")
    p_info.style = document.styles["Heading 2"]
    p_info.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p_info.runs:
        r.font.color.rgb = strong_blue

    table_info = document.add_table(rows=5, cols=2)
    table_info.style = "Table Grid"

    labels = [
        "First Name:",
        "Last Name:",
        "ID / Identification number:",
        "Email:",
        "Observation:",
    ]
    for i, label in enumerate(labels):
        row = table_info.rows[i]
        cell_label = row.cells[0]
        cell_val = row.cells[1]

        r_label = cell_label.paragraphs[0].add_run(label)
        r_label.bold = True
        r_label.font.name = "Calibri"
        r_label.font.size = Pt(10)

        cell_val.paragraphs[0].add_run("")

    _no_row_split(table_info)
    document.add_paragraph()

    # ------------------------------------------------------------------
    # 2) Roles
    # ------------------------------------------------------------------
    p_roles = document.add_paragraph("Roles (mark with X the requested roles)")
    p_roles.style = document.styles["Heading 2"]
    p_roles.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p_roles.runs:
        r.font.color.rgb = strong_blue

    table_roles = document.add_table(rows=1, cols=2)
    table_roles.style = "Table Grid"

    hdr_row = table_roles.rows[0]
    hdr_row.cells[0].paragraphs[0].add_run("X").bold = True
    hdr_row.cells[1].paragraphs[0].add_run("Role").bold = True

    # 👉 ahora usamos get_role_label_en para mostrar EN INGLÉS
    for rol in Rol.objects.all().order_by("nombre"):
        row = table_roles.add_row()
        row.cells[0].paragraphs[0].add_run("")  # X vacío
        label_en = get_role_label_en(rol.nombre)
        txt = row.cells[1].paragraphs[0].add_run(label_en)
        txt.font.name = "Calibri"
        txt.font.size = Pt(10)

    _no_row_split(table_roles)
    document.add_paragraph()

    # ------------------------------------------------------------------
    # 3) Projects
    # ------------------------------------------------------------------
    p_proy = document.add_paragraph(
        "Projects (mark with X the projects this user should access)"
    )
    p_proy.style = document.styles["Heading 2"]
    p_proy.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for r in p_proy.runs:
        r.font.color.rgb = strong_blue

    table_proj = document.add_table(rows=1, cols=7)
    table_proj.style = "Table Grid"

    headers = ["X", "Code", "Project name", "Client", "City", "State", "Office"]
    for idx, text in enumerate(headers):
        cell = table_proj.rows[0].cells[idx]
        run = cell.paragraphs[0].add_run(text)
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(10)

    try:
        proyectos_qs = Proyecto.objects.all().order_by("id")
    except Exception:
        proyectos_qs = []

    for p in proyectos_qs:
        row = table_proj.add_row()
        row.cells[0].paragraphs[0].add_run("")  # X
        row.cells[1].paragraphs[0].add_run(getattr(p, "codigo", "") or "")
        row.cells[2].paragraphs[0].add_run(getattr(p, "nombre", "") or "")
        row.cells[3].paragraphs[0].add_run(getattr(p, "mandante", "") or "")
        row.cells[4].paragraphs[0].add_run(getattr(p, "ciudad", "") or "")
        row.cells[5].paragraphs[0].add_run(getattr(p, "estado", "") or "")
        row.cells[6].paragraphs[0].add_run(getattr(p, "oficina", "") or "")

    _no_row_split(table_proj)

    # ------------------------------------------------------------------
    # Respuesta HTTP
    # ------------------------------------------------------------------
    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-"
            "officedocument.wordprocessingml.document"
        ),
    )
    response["Content-Disposition"] = (
        'attachment; filename="user_creation_request.docx"'
    )
    return response


@login_required(login_url="usuarios:login")
@rol_requerido("admin")
def gestionar_asignaciones_proyectos_view(request):
    from django.db import transaction
    from django.utils import timezone

    from facturacion.models import Proyecto
    from usuarios.models import ProyectoAsignacion

    # ==============================================================
    # SEGURIDAD
    # ==============================================================

    actor_is_admin = (
        request.user.is_superuser
        or request.user.roles.filter(nombre__iexact="admin").exists()
    )

    if not actor_is_admin:
        messages.error(
            request,
            "You do not have permission to manage project assignments.",
        )
        return redirect("dashboard_admin:listar_usuarios")

    # ==============================================================
    # SOLO PROYECTOS ACTIVOS
    # ==============================================================

    proyectos = Proyecto.objects.filter(activo=True).order_by(
        "nombre",
        "codigo",
    )

    # ==============================================================
    # PROYECTO SELECCIONADO
    # ==============================================================

    proyecto_id = (
        request.POST.get("proyecto_id")
        if request.method == "POST"
        else request.GET.get("proyecto")
    )

    proyecto_seleccionado = None

    if proyecto_id:
        try:
            proyecto_seleccionado = Proyecto.objects.get(
                id=proyecto_id,
                activo=True,
            )
        except (
            Proyecto.DoesNotExist,
            ValueError,
            TypeError,
        ):
            messages.error(
                request,
                "The selected project is invalid or inactive.",
            )

            return redirect("dashboard_admin:gestionar_asignaciones_proyectos")

    # ==============================================================
    # GUARDAR ASIGNACIONES
    # ==============================================================

    if request.method == "POST" and "save_project_assignments" in request.POST:
        if not proyecto_seleccionado:
            messages.error(
                request,
                "Please select an active project.",
            )

            return redirect("dashboard_admin:gestionar_asignaciones_proyectos")

        # ==========================================================
        # PALABRA DE SEGURIDAD
        # ==========================================================

        confirmation_word = request.POST.get("management_confirmation_word") or ""

        if not _valid_user_management_confirmation(confirmation_word):
            messages.error(
                request,
                "Invalid security confirmation word. "
                "No project assignments were changed.",
            )

            return redirect(
                "{}?proyecto={}".format(
                    reverse("dashboard_admin:gestionar_asignaciones_proyectos"),
                    proyecto_seleccionado.id,
                )
            )

        # ==========================================================
        # SOLO USUARIOS ACTIVOS SON OPERABLES
        # ==========================================================

        active_user_ids = set(
            User.objects.filter(is_active=True).values_list(
                "id",
                flat=True,
            )
        )

        requested_user_ids_raw = request.POST.getlist("usuarios")

        try:
            requested_user_ids = {int(user_id) for user_id in requested_user_ids_raw}

        except (
            TypeError,
            ValueError,
        ):
            messages.error(
                request,
                "One or more selected users are invalid.",
            )

            return redirect(
                "{}?proyecto={}".format(
                    reverse("dashboard_admin:gestionar_asignaciones_proyectos"),
                    proyecto_seleccionado.id,
                )
            )

        # No permitir IDs inactivos o inexistentes vía POST.
        if not requested_user_ids.issubset(active_user_ids):
            messages.error(
                request,
                "One or more selected users are invalid or inactive.",
            )

            return redirect(
                "{}?proyecto={}".format(
                    reverse("dashboard_admin:gestionar_asignaciones_proyectos"),
                    proyecto_seleccionado.id,
                )
            )

        # ==========================================================
        # VISIBILIDAD PARA NUEVAS ASIGNACIONES
        # ==========================================================

        visibility_mode = (request.POST.get("project_visibility") or "history").strip()

        if visibility_mode not in {
            "history",
            "from_now",
        }:
            messages.error(
                request,
                "Invalid project visibility mode.",
            )

            return redirect(
                "{}?proyecto={}".format(
                    reverse("dashboard_admin:gestionar_asignaciones_proyectos"),
                    proyecto_seleccionado.id,
                )
            )

        start_date_str = (request.POST.get("project_start_date") or "").strip()

        start_dt = None

        if visibility_mode == "from_now":
            try:
                if start_date_str:
                    parsed_date = timezone.datetime.fromisoformat(start_date_str)

                    if timezone.is_naive(parsed_date):
                        start_dt = timezone.make_aware(parsed_date)
                    else:
                        start_dt = parsed_date

                else:
                    start_dt = timezone.now()

            except (
                TypeError,
                ValueError,
            ):
                messages.error(
                    request,
                    "The start date is invalid.",
                )

                return redirect(
                    "{}?proyecto={}".format(
                        reverse("dashboard_admin:gestionar_asignaciones_proyectos"),
                        proyecto_seleccionado.id,
                    )
                )

        include_history = visibility_mode == "history"

        # ==========================================================
        # ESTADO ACTUAL
        #
        # Solamente usuarios ACTIVOS participan en altas/bajas.
        # Los usuarios inactivos continúan intactos.
        # ==========================================================

        current_active_assignments = list(
            ProyectoAsignacion.objects.filter(
                proyecto=proyecto_seleccionado,
                usuario__is_active=True,
            ).select_related("usuario")
        )

        current_active_ids = {
            assignment.usuario_id for assignment in current_active_assignments
        }

        current_assignments_by_user = {
            assignment.usuario_id: assignment
            for assignment in current_active_assignments
        }

        ids_to_add = requested_user_ids - current_active_ids

        ids_to_remove = current_active_ids - requested_user_ids

        ids_unchanged = current_active_ids & requested_user_ids

        # ==========================================================
        # CAMBIOS DE VISIBILIDAD EN ASIGNACIONES EXISTENTES
        #
        # Cada usuario ya asignado puede cambiar individualmente:
        #
        # history  -> include_history=True / start_at=None
        # from_now -> include_history=False / start_at=<fecha>
        #
        # Si permanece seleccionado pero no cambia esta configuración,
        # la asignación queda intacta.
        # ==========================================================

        visibility_updates = []

        for user_id in sorted(ids_unchanged):
            assignment = current_assignments_by_user.get(user_id)

            if not assignment:
                continue

            field_name = f"existing_visibility_{user_id}"

            requested_existing_mode = (request.POST.get(field_name) or "").strip()

            # Si el formulario no envió configuración para este usuario,
            # se conserva exactamente como estaba.
            if not requested_existing_mode:
                continue

            if requested_existing_mode not in {
                "history",
                "from_now",
            }:
                messages.error(
                    request,
                    "Invalid visibility setting for an existing assignment.",
                )

                return redirect(
                    "{}?proyecto={}".format(
                        reverse("dashboard_admin:gestionar_asignaciones_proyectos"),
                        proyecto_seleccionado.id,
                    )
                )

            if requested_existing_mode == "history":
                new_include_history = True
                new_start_at = None

            else:
                new_include_history = False

                existing_date_str = (
                    request.POST.get(f"existing_start_date_{user_id}") or ""
                ).strip()

                try:
                    if existing_date_str:
                        parsed_date = timezone.datetime.fromisoformat(existing_date_str)

                        if timezone.is_naive(parsed_date):
                            new_start_at = timezone.make_aware(parsed_date)
                        else:
                            new_start_at = parsed_date

                    elif not assignment.include_history and assignment.start_at:
                        # Si ya era "From date" y no se cambió la fecha,
                        # conserva la fecha existente.
                        new_start_at = assignment.start_at

                    else:
                        messages.error(
                            request,
                            "Please select a start date for " f"{assignment.usuario}.",
                        )

                        return redirect(
                            "{}?proyecto={}".format(
                                reverse(
                                    "dashboard_admin:"
                                    "gestionar_asignaciones_proyectos"
                                ),
                                proyecto_seleccionado.id,
                            )
                        )

                except (
                    TypeError,
                    ValueError,
                ):
                    messages.error(
                        request,
                        "One of the assignment start dates is invalid.",
                    )

                    return redirect(
                        "{}?proyecto={}".format(
                            reverse(
                                "dashboard_admin:" "gestionar_asignaciones_proyectos"
                            ),
                            proyecto_seleccionado.id,
                        )
                    )

            current_include_history = assignment.include_history

            current_start_date = (
                assignment.start_at.date() if assignment.start_at else None
            )

            new_start_date = new_start_at.date() if new_start_at else None

            if (
                current_include_history != new_include_history
                or current_start_date != new_start_date
            ):
                visibility_updates.append(
                    (
                        assignment,
                        new_include_history,
                        new_start_at,
                    )
                )

        # ==========================================================
        # ESCRITURA ATÓMICA
        # ==========================================================

        with transaction.atomic():

            # ------------------------------------------------------
            # QUITAR
            #
            # SOLO usuarios activos desmarcados.
            # Usuarios inactivos NO se tocan.
            # ------------------------------------------------------

            if ids_to_remove:
                ProyectoAsignacion.objects.filter(
                    proyecto=proyecto_seleccionado,
                    usuario_id__in=ids_to_remove,
                    usuario__is_active=True,
                ).delete()

            # ------------------------------------------------------
            # AGREGAR
            #
            # Usa la configuración global:
            # "Visibility for new assignments".
            # ------------------------------------------------------

            nuevos = []

            for user_id in sorted(ids_to_add):
                nuevos.append(
                    ProyectoAsignacion(
                        usuario_id=user_id,
                        proyecto=proyecto_seleccionado,
                        include_history=include_history,
                        start_at=(
                            None if include_history else (start_dt or timezone.now())
                        ),
                    )
                )

            if nuevos:
                ProyectoAsignacion.objects.bulk_create(nuevos)

            # ------------------------------------------------------
            # MODIFICAR VISIBILIDAD DE ASIGNACIONES EXISTENTES
            #
            # No elimina ni recrea la asignación.
            # Solamente modifica include_history/start_at.
            # ------------------------------------------------------

            for (
                assignment,
                new_include_history,
                new_start_at,
            ) in visibility_updates:

                assignment.include_history = new_include_history

                assignment.start_at = new_start_at

                assignment.save(
                    update_fields=[
                        "include_history",
                        "start_at",
                    ]
                )

        messages.success(
            request,
            (
                f'Project "{proyecto_seleccionado.nombre}" '
                f"updated successfully. "
                f"Added: {len(ids_to_add)}. "
                f"Removed: {len(ids_to_remove)}. "
                f"Visibility changed: {len(visibility_updates)}. "
                f"Unchanged: "
                f"{len(ids_unchanged) - len(visibility_updates)}."
            ),
        )

        return redirect(
            "{}?proyecto={}".format(
                reverse("dashboard_admin:gestionar_asignaciones_proyectos"),
                proyecto_seleccionado.id,
            )
        )

    # ==============================================================
    # USUARIOS ACTIVOS
    # ==============================================================

    usuarios = (
        User.objects.filter(is_active=True)
        .prefetch_related("roles")
        .order_by(
            "first_name",
            "last_name",
            "username",
        )
    )

    # ==============================================================
    # ASIGNACIONES ACTUALES DEL PROYECTO
    # ==============================================================

    asignaciones_por_usuario = {}

    usuarios_asignados_ids = set()

    if proyecto_seleccionado:
        asignaciones_actuales = ProyectoAsignacion.objects.filter(
            proyecto=proyecto_seleccionado,
            usuario__is_active=True,
        ).select_related("usuario")

        for asignacion in asignaciones_actuales:
            usuarios_asignados_ids.add(asignacion.usuario_id)

            asignaciones_por_usuario[asignacion.usuario_id] = asignacion

    # ==============================================================
    # CONSTRUIR FILAS PARA TEMPLATE
    # ==============================================================

    usuarios_rows = []

    for usuario in usuarios:
        asignacion = asignaciones_por_usuario.get(usuario.id)

        usuarios_rows.append(
            {
                "usuario": usuario,
                "asignado": (usuario.id in usuarios_asignados_ids),
                "asignacion": asignacion,
            }
        )

    # ==============================================================
    # CONTADORES
    # ==============================================================

    total_usuarios_activos = len(usuarios_rows)

    total_asignados_activos = len(usuarios_asignados_ids)

    # Solo auditoría informativa.
    # No se muestran ni se manipulan.
    total_asignados_inactivos = 0

    if proyecto_seleccionado:
        total_asignados_inactivos = ProyectoAsignacion.objects.filter(
            proyecto=proyecto_seleccionado,
            usuario__is_active=False,
        ).count()

    return render(
        request,
        "dashboard_admin/gestionar_asignaciones_proyectos.html",
        {
            "proyectos": proyectos,
            "proyecto_seleccionado": proyecto_seleccionado,
            "usuarios_rows": usuarios_rows,
            "total_usuarios_activos": total_usuarios_activos,
            "total_asignados_activos": total_asignados_activos,
            "total_asignados_inactivos": total_asignados_inactivos,
        },
    )
