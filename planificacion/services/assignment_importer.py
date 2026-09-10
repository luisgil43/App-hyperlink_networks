import io
import re
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from lxml import etree
from pypdf import PdfReader

SUPPORTED_EXTENSIONS = {
    ".pdf",
    ".docx",
}


DATE_PATTERNS = (
    "%m/%d/%Y",
    "%m/%d/%y",
    "%m-%d-%Y",
    "%m-%d-%y",
)


WORD_NAMESPACE = {
    "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
}


class AssignmentImportError(Exception):
    pass


def _clean_text(value):
    if value is None:
        return ""

    value = str(value)

    value = value.replace(
        "\u00a0",
        " ",
    )

    value = value.replace(
        "\u2013",
        "-",
    )

    value = value.replace(
        "\u2014",
        "-",
    )

    value = value.replace(
        "\u2019",
        "'",
    )

    value = value.replace(
        "\u2032",
        "'",
    )

    value = re.sub(
        r"[ \t]+",
        " ",
        value,
    )

    value = re.sub(
        r"\s*\n\s*",
        " ",
        value,
    )

    return value.strip()


def _normalize_lines(text):
    result = []

    for raw_line in text.splitlines():
        line = _clean_text(raw_line)

        if line:
            result.append(line)

    return result


def _normalize_code(value):
    value = _clean_text(value).upper()

    value = re.sub(
        r"\s+",
        "",
        value,
    )

    return value


def _read_pdf(uploaded_file):
    try:
        uploaded_file.seek(0)

        reader = PdfReader(uploaded_file)

        pages = []

        for page in reader.pages:
            text = page.extract_text() or ""

            if text.strip():
                pages.append(text)

        return {
            "type": "pdf",
            "text": "\n".join(pages),
            "xml_lines": [],
            "headers": [],
            "tables": [],
        }

    except Exception as exc:
        raise AssignmentImportError(f"Unable to read PDF: {exc}") from exc


def _extract_docx_headers(document):
    headers = []

    for section in document.sections:
        header = section.header

        for paragraph in header.paragraphs:
            text = _clean_text(paragraph.text)

            if text and text not in headers:
                headers.append(text)

        for table in header.tables:
            for row in table.rows:
                values = [_clean_text(cell.text) for cell in row.cells]

                values = [value for value in values if value]

                if values:
                    joined = " | ".join(values)

                    if joined not in headers:
                        headers.append(joined)

    return headers


def _extract_docx_tables(document):
    tables = []

    for table in document.tables:
        rows = []

        for row in table.rows:
            values = [_clean_text(cell.text) for cell in row.cells]

            rows.append(values)

        if rows:
            tables.append(rows)

    return tables


def _extract_docx_xml_lines(raw):
    """
    Read Word's underlying document XML.

    Some client NTP documents contain visible text in structures that
    python-docx does not expose as normal table cells.

    Joining all w:t nodes belonging to the same w:p also repairs values
    split internally by Word, for example:

        0913TA_ + 04
        09/ + 1 + 8 + /2026
        C-10 + 7 + " FIBER PLACEMENT"
    """

    try:
        with ZipFile(io.BytesIO(raw)) as archive:
            xml = archive.read("word/document.xml")

    except Exception as exc:
        raise AssignmentImportError(f"Unable to inspect DOCX XML: {exc}") from exc

    try:
        root = etree.fromstring(xml)

    except Exception as exc:
        raise AssignmentImportError(f"Unable to parse DOCX XML: {exc}") from exc

    lines = []

    paragraphs = root.xpath(
        "//w:p",
        namespaces=WORD_NAMESPACE,
    )

    for paragraph in paragraphs:
        fragments = paragraph.xpath(
            ".//w:t/text()",
            namespaces=WORD_NAMESPACE,
        )

        if not fragments:
            continue

        value = "".join(fragments)

        value = _clean_text(value)

        if value:
            lines.append(value)

    return lines


def _read_docx(uploaded_file):
    try:
        uploaded_file.seek(0)

        raw = uploaded_file.read()

        document = Document(io.BytesIO(raw))

        headers = _extract_docx_headers(document)

        tables = _extract_docx_tables(document)

        xml_lines = _extract_docx_xml_lines(raw)

        blocks = []

        blocks.extend(headers)

        blocks.extend(xml_lines)

        for paragraph in document.paragraphs:
            text = _clean_text(paragraph.text)

            if text and text not in blocks:
                blocks.append(text)

        for table in tables:
            for row in table:
                values = [value for value in row if value]

                if values:
                    joined = " | ".join(values)

                    if joined not in blocks:
                        blocks.append(joined)

        return {
            "type": "docx",
            "text": "\n".join(blocks),
            "xml_lines": xml_lines,
            "headers": headers,
            "tables": tables,
        }

    except AssignmentImportError:
        raise

    except Exception as exc:
        raise AssignmentImportError(f"Unable to read DOCX: {exc}") from exc


def read_document(uploaded_file):
    if uploaded_file is None:
        raise AssignmentImportError("Select a PDF or DOCX file first.")

    filename = uploaded_file.name or ""

    extension = Path(filename).suffix.lower()

    if extension not in SUPPORTED_EXTENSIONS:
        raise AssignmentImportError("Unsupported file type. Use PDF or DOCX.")

    if extension == ".pdf":
        document_data = _read_pdf(uploaded_file)

    else:
        document_data = _read_docx(uploaded_file)

    text = document_data["text"].strip()

    if not text:
        raise AssignmentImportError(
            "The document does not contain extractable text. "
            "It may be a scanned PDF and will require OCR."
        )

    return document_data


def _parse_date(value):
    value = _clean_text(value)

    value = re.sub(
        r"\s+",
        "",
        value,
    )

    for pattern in DATE_PATTERNS:
        try:
            return datetime.strptime(
                value,
                pattern,
            ).date()

        except ValueError:
            continue

    return None


def _parse_decimal(raw_value):
    raw_value = _clean_text(raw_value)

    raw_value = raw_value.replace(
        ",",
        "",
    )

    match = re.search(
        r"-?\d+(?:\.\d+)?",
        raw_value,
    )

    if not match:
        return None

    try:
        return Decimal(match.group(0))

    except (
        InvalidOperation,
        ValueError,
    ):
        return None


def _find_table_by_headers(
    tables,
    required_headers,
):
    required_headers = {_clean_text(header).upper() for header in required_headers}

    for table in tables:
        if not table:
            continue

        first_row = {
            _clean_text(value).upper() for value in table[0] if _clean_text(value)
        }

        if required_headers.issubset(first_row):
            return table

    return None


def _table_column_indexes(header_row):
    result = {}

    for index, value in enumerate(header_row):
        key = _clean_text(value).upper()

        if key:
            result[key] = index

    return result


def _cell(
    row,
    indexes,
    column_name,
):
    index = indexes.get(column_name.upper())

    if index is None:
        return ""

    if index >= len(row):
        return ""

    return _clean_text(row[index])


def _find_sequence_after_headers(
    lines,
    header_names,
):
    """
    Find a consecutive logical block such as:

        ISSUE DATE
        PC
        MARKET
        CONTRACTOR
        DFN
        09/08/2026
        676
        GREEN BAY
        HYPERLINK
        0913TA_04
    """

    normalized_headers = [_clean_text(value).upper() for value in header_names]

    upper_lines = [_clean_text(value).upper() for value in lines]

    for start_index in range(len(upper_lines)):
        indexes = []

        cursor = start_index

        valid = True

        for expected_header in normalized_headers:
            found_index = None

            for index in range(
                cursor,
                min(
                    cursor + 4,
                    len(upper_lines),
                ),
            ):
                if upper_lines[index] == expected_header:
                    found_index = index
                    break

            if found_index is None:
                valid = False
                break

            indexes.append(found_index)

            cursor = found_index + 1

        if not valid:
            continue

        last_header_index = indexes[-1]

        values = lines[
            last_header_index + 1 : last_header_index + 1 + len(header_names)
        ]

        if len(values) == len(header_names):
            return {
                header_names[index]: _clean_text(values[index])
                for index in range(len(header_names))
            }

    return {}


def _extract_summary_from_xml_lines(
    xml_lines,
):
    result = {
        "issue_date": None,
        "pc": "",
        "market": "",
        "contractor": "",
        "dfn_codes": [],
    }

    values = _find_sequence_after_headers(
        xml_lines,
        [
            "ISSUE DATE",
            "PC",
            "MARKET",
            "CONTRACTOR",
            "DFN",
        ],
    )

    if not values:
        return result

    result["issue_date"] = _parse_date(
        values.get(
            "ISSUE DATE",
            "",
        )
    )

    result["pc"] = values.get(
        "PC",
        "",
    )

    result["market"] = values.get(
        "MARKET",
        "",
    )

    result["contractor"] = values.get(
        "CONTRACTOR",
        "",
    )

    dfn_code = _normalize_code(
        values.get(
            "DFN",
            "",
        )
    )

    if dfn_code:
        result["dfn_codes"].append(dfn_code)

    return result


def _extract_summary_from_docx_tables(
    tables,
):
    result = {
        "issue_date": None,
        "pc": "",
        "market": "",
        "contractor": "",
        "dfn_codes": [],
    }

    table = _find_table_by_headers(
        tables,
        {
            "ISSUE DATE",
            "PC",
            "MARKET",
            "CONTRACTOR",
            "DFN",
        },
    )

    if table is None or len(table) < 2:
        return result

    indexes = _table_column_indexes(table[0])

    row = table[1]

    result["issue_date"] = _parse_date(
        _cell(
            row,
            indexes,
            "ISSUE DATE",
        )
    )

    result["pc"] = _cell(
        row,
        indexes,
        "PC",
    )

    result["market"] = _cell(
        row,
        indexes,
        "MARKET",
    )

    result["contractor"] = _cell(
        row,
        indexes,
        "CONTRACTOR",
    )

    dfn_code = _normalize_code(
        _cell(
            row,
            indexes,
            "DFN",
        )
    )

    if dfn_code:
        result["dfn_codes"].append(dfn_code)

    return result


def _merge_summary(
    primary,
    secondary,
):
    result = {
        "issue_date": (primary.get("issue_date") or secondary.get("issue_date")),
        "pc": (primary.get("pc") or secondary.get("pc") or ""),
        "market": (primary.get("market") or secondary.get("market") or ""),
        "contractor": (primary.get("contractor") or secondary.get("contractor") or ""),
        "dfn_codes": [],
    }

    for source in (
        primary,
        secondary,
    ):
        for code in source.get(
            "dfn_codes",
            [],
        ):
            normalized = _normalize_code(code)

            if normalized and normalized not in result["dfn_codes"]:
                result["dfn_codes"].append(normalized)

    return result


def _extract_dfns_from_text(
    text,
):
    results = []

    patterns = (
        r"\b\d{3,6}[A-Z]{1,4}_[A-Z0-9]{1,10}\b",
        r"\b\d{3,6}[A-Z]{1,4}\s*_\s*[A-Z0-9]{1,10}\b",
    )

    for pattern in patterns:
        for match in re.finditer(
            pattern,
            text,
            re.IGNORECASE,
        ):
            code = _normalize_code(match.group(0))

            if code and code not in results:
                results.append(code)

    return results


def _extract_schedule_activities_from_tables(
    tables,
):
    activities = {}

    table = _find_table_by_headers(
        tables,
        {
            "TASK",
            "START DATE",
            "END DATE",
        },
    )

    if table is None:
        return activities

    indexes = _table_column_indexes(table[0])

    for row in table[1:]:
        task = _cell(
            row,
            indexes,
            "TASK",
        )

        code_match = re.search(
            r"\bC-\d{2,5}\b",
            task,
            re.IGNORECASE,
        )

        if not code_match:
            continue

        code = code_match.group(0).upper()

        name = re.sub(
            re.escape(code_match.group(0)),
            "",
            task,
            count=1,
            flags=re.IGNORECASE,
        )

        name = _clean_text(name)

        activities[code] = {
            "code": code,
            "name": name,
            "quantity": None,
            "unit": "",
            "client_start_date": _parse_date(
                _cell(
                    row,
                    indexes,
                    "START DATE",
                )
            ),
            "client_end_date": _parse_date(
                _cell(
                    row,
                    indexes,
                    "END DATE",
                )
            ),
        }

    return activities


def _extract_schedule_activities_from_xml(
    xml_lines,
):
    activities = {}

    upper_lines = [value.upper() for value in xml_lines]

    try:
        header_index = next(
            index
            for index in range(len(upper_lines) - 2)
            if (
                upper_lines[index] == "TASK"
                and upper_lines[index + 1] == "START DATE"
                and upper_lines[index + 2] == "END DATE"
            )
        )

    except StopIteration:
        return activities

    cursor = header_index + 3

    while cursor < len(xml_lines):
        line = xml_lines[cursor]

        if line.upper() == "TASK":
            break

        code_match = re.search(
            r"\bC-\d{2,5}\b",
            line,
            re.IGNORECASE,
        )

        if not code_match:
            cursor += 1
            continue

        code = code_match.group(0).upper()

        name = re.sub(
            re.escape(code_match.group(0)),
            "",
            line,
            count=1,
            flags=re.IGNORECASE,
        )

        name = _clean_text(name)

        start_date = None
        end_date = None

        if cursor + 1 < len(xml_lines):
            start_date = _parse_date(xml_lines[cursor + 1])

        if cursor + 2 < len(xml_lines):
            end_date = _parse_date(xml_lines[cursor + 2])

        activities[code] = {
            "code": code,
            "name": name,
            "quantity": None,
            "unit": "",
            "client_start_date": start_date,
            "client_end_date": end_date,
        }

        cursor += 3

    return activities


def _infer_unit_from_activity(
    code,
    name,
    quantity_text,
):
    upper = (f"{code} {name} {quantity_text}").upper()

    if (
        "'" in quantity_text
        or "FT" in upper
        or "FEET" in upper
        or "FOOT" in upper
        or "FIBER PLACEMENT" in upper
    ):
        return "ft"

    if any(
        value in upper
        for value in (
            "SPLIC",
            "TEST",
            "BOX",
            "CTO",
        )
    ):
        return "box"

    return ""


def _extract_quantity_activities_from_tables(
    tables,
):
    result = {}

    table = _find_table_by_headers(
        tables,
        {
            "TASK",
            "QUANTITY",
        },
    )

    if table is None:
        return result

    indexes = _table_column_indexes(table[0])

    for row in table[1:]:
        task = _cell(
            row,
            indexes,
            "TASK",
        )

        quantity_text = _cell(
            row,
            indexes,
            "QUANTITY",
        )

        code_match = re.search(
            r"\bC-\d{2,5}\b",
            task,
            re.IGNORECASE,
        )

        if not code_match:
            continue

        code = code_match.group(0).upper()

        name = re.sub(
            re.escape(code_match.group(0)),
            "",
            task,
            count=1,
            flags=re.IGNORECASE,
        )

        name = _clean_text(name)

        cleaned_quantity_text = re.sub(
            r"\bC-\d{2,5}\b",
            "",
            quantity_text,
            count=1,
            flags=re.IGNORECASE,
        )

        cleaned_quantity_text = re.sub(
            r"\*ESTIMATED\b",
            "",
            cleaned_quantity_text,
            flags=re.IGNORECASE,
        )

        quantity = _parse_decimal(cleaned_quantity_text)

        unit = _infer_unit_from_activity(
            code,
            name,
            cleaned_quantity_text,
        )

        result[code] = {
            "code": code,
            "name": name,
            "quantity": quantity,
            "unit": unit,
        }

    return result


def _extract_quantity_activities_from_xml(
    xml_lines,
):
    result = {}

    upper_lines = [value.upper() for value in xml_lines]

    quantity_header_index = None

    for index in range(len(upper_lines) - 1):
        if upper_lines[index] == "TASK" and upper_lines[index + 1] == "QUANTITY":
            quantity_header_index = index
            break

    if quantity_header_index is None:
        return result

    cursor = quantity_header_index + 4

    while cursor < len(xml_lines):
        task = xml_lines[cursor]

        code_match = re.search(
            r"\bC-\d{2,5}\b",
            task,
            re.IGNORECASE,
        )

        if not code_match:
            cursor += 1
            continue

        code = code_match.group(0).upper()

        name = re.sub(
            re.escape(code_match.group(0)),
            "",
            task,
            count=1,
            flags=re.IGNORECASE,
        )

        name = _clean_text(name)

        quantity = None
        unit = ""

        lookahead = xml_lines[
            cursor
            + 1 : min(
                cursor + 5,
                len(xml_lines),
            )
        ]

        for value in lookahead:
            if re.match(r"^\$", value):
                break

            cleaned = re.sub(
                r"\bC-\d{2,5}\b",
                "",
                value,
                flags=re.IGNORECASE,
            )

            cleaned = re.sub(
                r"\*ESTIMATED\b",
                "",
                cleaned,
                flags=re.IGNORECASE,
            )

            parsed = _parse_decimal(cleaned)

            if parsed is None:
                continue

            quantity = parsed

            unit = _infer_unit_from_activity(
                code,
                name,
                cleaned,
            )

            break

        if quantity is not None and not unit:
            unit = _infer_unit_from_activity(
                code,
                name,
                "",
            )

        result[code] = {
            "code": code,
            "name": name,
            "quantity": quantity,
            "unit": unit,
        }

        cursor += 1

    return result


def _merge_activity_sources(
    schedule_primary,
    schedule_secondary,
    quantity_primary,
    quantity_secondary,
):
    codes = []

    for source in (
        schedule_primary,
        schedule_secondary,
        quantity_primary,
        quantity_secondary,
    ):
        for code in source:
            if code not in codes:
                codes.append(code)

    result = []

    for code in codes:
        schedule = schedule_primary.get(code) or schedule_secondary.get(code) or {}

        quantity = quantity_primary.get(code) or quantity_secondary.get(code) or {}

        name = schedule.get("name") or quantity.get("name") or ""

        result.append(
            {
                "code": code,
                "name": name,
                "quantity": quantity.get("quantity"),
                "unit": quantity.get(
                    "unit",
                    "",
                ),
                "client_start_date": schedule.get("client_start_date"),
                "client_end_date": schedule.get("client_end_date"),
            }
        )

    return result


def _extract_generic_activity_codes(
    text,
):
    codes = []

    for match in re.finditer(
        r"\bC-\d{2,5}\b",
        text,
        re.IGNORECASE,
    ):
        code = match.group(0).upper()

        if code not in codes:
            codes.append(code)

    return codes


def _generic_pdf_activities(
    text,
):
    lines = _normalize_lines(text)

    activities = []

    for code in _extract_generic_activity_codes(text):
        matching_lines = [line for line in lines if code in line.upper()]

        surrounding_text = " ".join(matching_lines)

        dates = re.findall(
            r"\b\d{1,2}[/-]\d{1,2}[/-]\d{2,4}\b",
            surrounding_text,
        )

        parsed_dates = [_parse_date(value) for value in dates]

        parsed_dates = [value for value in parsed_dates if value is not None]

        quantity = None
        unit = ""

        foot_match = re.search(
            r"(\d[\d,]*(?:\.\d+)?)\s*(?:FT|FEET|FOOT|'|′)",
            surrounding_text,
            re.IGNORECASE,
        )

        if foot_match:
            quantity = _parse_decimal(foot_match.group(1))

            unit = "ft"

        activities.append(
            {
                "code": code,
                "name": "",
                "quantity": quantity,
                "unit": unit,
                "client_start_date": (
                    parsed_dates[0] if len(parsed_dates) >= 1 else None
                ),
                "client_end_date": (
                    parsed_dates[1] if len(parsed_dates) >= 2 else None
                ),
            }
        )

    return activities


def _extract_client(
    text,
    filename,
):
    upper_text = text.upper()

    client_patterns = (
        r"\b(ITG)\s+RESERVES\b",
        r"\b(ITG)\s+WISCONSIN\b",
    )

    for pattern in client_patterns:
        match = re.search(
            pattern,
            upper_text,
            re.IGNORECASE,
        )

        if match:
            return match.group(1).upper()

    upper_filename = filename.upper()

    if re.search(
        r"\bITG\b",
        upper_filename,
    ):
        return "ITG"

    return ""


def _extract_assignment_type(
    text,
    filename,
):
    combined = (text + "\n" + filename).upper()

    if "NOTICE TO PROCEED" in combined or re.search(
        r"\bNTP\b",
        combined,
    ):
        return "NTP"

    return ""


def extract_assignment(uploaded_file):
    document_data = read_document(uploaded_file)

    text = document_data["text"]

    xml_lines = document_data.get(
        "xml_lines",
        [],
    )

    tables = document_data.get(
        "tables",
        [],
    )

    document_type = document_data.get("type")

    filename = uploaded_file.name or "Imported Assignment"

    if document_type == "docx":
        xml_summary = _extract_summary_from_xml_lines(xml_lines)

        table_summary = _extract_summary_from_docx_tables(tables)

        summary = _merge_summary(
            xml_summary,
            table_summary,
        )

        schedule_from_tables = _extract_schedule_activities_from_tables(tables)

        schedule_from_xml = _extract_schedule_activities_from_xml(xml_lines)

        quantity_from_tables = _extract_quantity_activities_from_tables(tables)

        quantity_from_xml = _extract_quantity_activities_from_xml(xml_lines)

        activities = _merge_activity_sources(
            schedule_from_tables,
            schedule_from_xml,
            quantity_from_tables,
            quantity_from_xml,
        )

    else:
        summary = {
            "issue_date": None,
            "pc": "",
            "market": "",
            "contractor": "",
            "dfn_codes": [],
        }

        activities = _generic_pdf_activities(text)

    dfn_codes = list(
        summary.get(
            "dfn_codes",
            [],
        )
    )

    for code in _extract_dfns_from_text(text):
        if code not in dfn_codes:
            dfn_codes.append(code)

    client_name = _extract_client(
        text,
        filename,
    )

    assignment_type = _extract_assignment_type(
        text,
        filename,
    )

    warnings = []

    if not client_name:
        warnings.append("Client was not detected.")

    if not summary.get("issue_date"):
        warnings.append("Issue Date was not detected.")

    if not summary.get("market"):
        warnings.append("Market was not detected.")

    if not summary.get("contractor"):
        warnings.append("Contractor was not detected.")

    if not dfn_codes:
        warnings.append("No DFN was detected.")

    if not activities:
        warnings.append("No activities were detected.")

    for activity in activities:
        missing = []

        if activity["quantity"] is None:
            missing.append("quantity")

        if not activity["unit"]:
            missing.append("unit")

        if activity["client_start_date"] is None:
            missing.append("Client Start")

        if activity["client_end_date"] is None:
            missing.append("Client Finish")

        if missing:
            warnings.append(f'{activity["code"]}: review ' + ", ".join(missing) + ".")

    assignment_name = Path(filename).stem

    assignment_name = re.sub(
        r"\(\d+\)$",
        "",
        assignment_name,
    )

    assignment_name = assignment_name.replace(
        "_",
        " ",
    ).strip()

    return {
        "filename": filename,
        "raw_text": text,
        "assignment": {
            "client_name": client_name,
            "name": assignment_name,
            "assignment_type": assignment_type,
            "client_reference": "",
            "issue_date": summary.get("issue_date"),
            "pc": summary.get(
                "pc",
                "",
            ),
            "contractor": summary.get(
                "contractor",
                "",
            ),
            "country": "United States",
            "state": "",
            "city": "",
            "market": summary.get(
                "market",
                "",
            ),
            "source_reference": (f"Imported from: {filename}"),
        },
        "dfns": [
            {
                "code": code,
            }
            for code in dfn_codes
        ],
        "activities": activities,
        "warnings": warnings,
    }
